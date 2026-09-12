import io
import re
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

router = APIRouter()

class TablePayload(BaseModel):
    raw_text: str

def clean_chatgpt_formatting(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'(\*\*|__|\*|_)(.*?)\1', r'\2', text)
    text = re.sub(r'(\*\*|__|\*|_)', '', text)
    return text.strip()

@router.post("/convert")
async def convert_table(payload: TablePayload):
    lines = [line.strip() for line in payload.raw_text.strip().splitlines() if line.strip()]
    if not lines:
        raise HTTPException(status_code=400, detail="Empty text provided.")

    rows = []
    for line in lines:
        if re.match(r'^[|:-_\s]+$', line):
            continue
        cells = line.split('\t') if '\t' in line else [c.strip() for c in line.split('|') if c.strip()]
        cleaned = [clean_chatgpt_formatting(c) for c in cells if c.strip()]
        if cleaned:
            rows.append(cleaned)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SOP"

    font_arial = Font(name="Arial", size=12)
    font_header = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    thin_black = Side(border_style="thin", color="000000")
    border_black = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)

    headers = ["No", "Work Procedure", "Quality Standard", "Working Point", "Tools/Material", "Phenomenon", "Safety"]
    spans = [1, 9, 4, 6, 3, 3, 4]

    # Render Header
    col_idx = 1
    for h_text, span in zip(headers, spans):
        start_col, end_col = col_idx, col_idx + span - 1
        ws.cell(row=1, column=start_col, value=h_text)
        if span > 1:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font, cell.border = header_fill, font_header, border_black
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_idx += span

    # Render Data
    current_row = 2
    for r_idx, r_data in enumerate(rows, start=1):
        content = [str(r_idx)] + r_data[:6]
        while len(content) < 7:
            content.append("")

        col_idx = 1
        for val, span in zip(content, spans):
            start_col, end_col = col_idx, col_idx + span - 1
            ws.cell(row=current_row, column=start_col, value=val)
            if span > 1:
                ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font, cell.border = font_arial, border_black
                cell.alignment = Alignment(horizontal="center" if span == 1 else "left", vertical="center", wrap_text=True)
            col_idx += span
        current_row += 1

    for c in range(1, 31):
        ws.column_dimensions[get_column_letter(c)].width = 6.0

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={'Content-Disposition': 'attachment; filename="Standardized_SOP.xlsx"'}
    )